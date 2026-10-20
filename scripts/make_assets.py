"""Builds the binary assets the app ships with: PNG icons and the blank
Excel template. Run from the repo root: python3 scripts/make_assets.py
Needs Pillow and openpyxl."""
from pathlib import Path

from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parent.parent
ICONS = ROOT / "app" / "icons"
TEAL = (15, 118, 110)
WHITE = (255, 255, 255)

# Column widths for the template, matching COLUMNS in app/js/fields.js.
HEADERS = [
    ("Captured At", 17), ("First Name", 14), ("Last Name", 16), ("Company", 24), ("Title", 24),
    ("Email", 28), ("Phone", 17), ("Tags", 20), ("Notes", 48), ("Event", 20), ("Captured By", 16),
    ("Source", 11), ("Other Phone", 17), ("Website", 24), ("LinkedIn", 28), ("Address", 32),
    ("Other Info", 28), ("Card Photo", 12), ("QR Data", 24), ("Updated At", 17), ("ID", 38),
]


def draw_icon(size, maskable=False, rounded=True):
    scale = 4  # draw large, downsample for smooth edges
    s = size * scale
    img = Image.new("RGBA", (s, s), (0, 0, 0, 0))
    d = ImageDraw.Draw(img)
    if maskable or not rounded:
        d.rectangle([0, 0, s, s], fill=TEAL)
    else:
        d.rounded_rectangle([0, 0, s - 1, s - 1], radius=int(s * 0.22), fill=TEAL)
    # Glyph: QR finder corners around a card. Maskable icons keep it inside
    # the central safe zone.
    pad = s * (0.27 if maskable else 0.2)
    span = s - 2 * pad
    w = max(2, int(s * 0.055))
    arm = span * 0.26
    x0, y0, x1, y1 = pad, pad, s - pad, s - pad
    for (cx, cy, dx, dy) in [(x0, y0, 1, 1), (x1, y0, -1, 1), (x0, y1, 1, -1), (x1, y1, -1, -1)]:
        d.line([(cx, cy + dy * arm), (cx, cy), (cx + dx * arm, cy)], fill=WHITE, width=w, joint="curve")
    # Card in the middle
    cw, ch = span * 0.62, span * 0.42
    cx0, cy0 = s / 2 - cw / 2, s / 2 - ch / 2
    d.rounded_rectangle([cx0, cy0, cx0 + cw, cy0 + ch], radius=int(ch * 0.16), fill=WHITE)
    r = ch * 0.17
    pcx, pcy = cx0 + cw * 0.26, cy0 + ch * 0.4
    d.ellipse([pcx - r, pcy - r, pcx + r, pcy + r], fill=TEAL)
    d.rounded_rectangle([pcx - r * 1.7, pcy + r * 1.35, pcx + r * 1.7, pcy + r * 2.9], radius=int(r), fill=TEAL)
    lx = cx0 + cw * 0.5
    lh = max(2, int(ch * 0.09))
    for i, frac in enumerate([0.32, 0.55]):
        ly = cy0 + ch * frac
        d.rounded_rectangle([lx, ly, lx + cw * (0.38 if i == 0 else 0.28), ly + lh], radius=lh // 2, fill=TEAL)
    return img.resize((size, size), Image.LANCZOS)


def make_icons():
    ICONS.mkdir(parents=True, exist_ok=True)
    draw_icon(192).save(ICONS / "icon-192.png")
    draw_icon(512).save(ICONS / "icon-512.png")
    draw_icon(512, maskable=True).save(ICONS / "icon-maskable-512.png")
    # iOS applies its own corner mask, so the touch icon is square.
    draw_icon(180, rounded=False).convert("RGB").save(ICONS / "apple-touch-icon.png")


def make_template():
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"  # left blank on purpose: the app adds the table with the first row
    for i, (_, width) in enumerate(HEADERS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.freeze_panes = "A2"
    about = wb.create_sheet("About")
    about["A1"] = "InfoCapture leads"
    about["A3"] = "Contacts captured with the InfoCapture app land in the table on the Leads sheet."
    about["A4"] = "You can sort, filter, and add your own columns. Keep the column headers as they are so the app can find them."
    about["A5"] = "Card photos are saved in the 'Card Photos' folder next to this workbook."
    about.column_dimensions["A"].width = 110
    out = ROOT / "app" / "template" / "InfoCapture-Leads.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)


if __name__ == "__main__":
    make_icons()
    make_template()
    print("icons and template written")
